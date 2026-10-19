"""
Lectura de la hoja "LISTA DE PRECIO" de los libros BASE DATOS <CIUDAD>.xlsm.

La hoja tiene una fila por servicio (ÍTEM, DESCRIPCIÓN, NORMA, PRECIO LISTA y,
en Bucaramanga e Ibagué, DESCUENTO) y después una columna por obra cuyo
encabezado es el código de obra del Excel ("8-1" = empresa 8, obra 1) con el
precio pactado para esa obra. Las columnas de obra se detectan por su
encabezado (Bogotá no tiene DESCUENTO y empiezan una columna antes).
Los códigos "0-N" son clientes varios: cada N es una empresa distinta, que
en el portal se llama BOG0-N (constructora y obra), con nombre en la hoja
LISTA EMPRESAS VARIOS OBRAS. Las filas cuyo ÍTEM no tiene guion y cuya DESCRIPCIÓN
va en mayúsculas son encabezados de categoría.

Limpieza aplicada (la misma que se usó para generar catalogo_bucaramanga.py):
- Separadores de código normalizados a guion ("2.11" -> "2-11", "11,1,1L" -> "11-1-1L").
- Código cuyo primer segmento no es el número de la categoría pero empieza por
  él ("13-8" dentro de CONCRETOS) se corrige reinsertando el guion ("1-3-8").
- Filas con el mismo nombre dentro de una categoría se fusionan (normas unidas
  con " / "; los precios de la fila descartada rellenan los huecos de la conservada).
- Nombres repetidos entre categorías reciben sufijo: "(Laboratorista)",
  "(Auxiliar)", "(Laboratorista mes)" o "(cód. X)".
- Códigos repetidos con distinto nombre reciben sufijo "-2", "-3", ...
- Norma "N/A" o vacía se guarda como cadena vacía. Precios <= 0 se ignoran.

Se usa desde `python manage.py importar_catalogo_excel`.
"""

import re
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal

from django.db.models import Q

from core.models import Obra
from core.wp_normalizar import misma_empresa
from users.ciudades import CIUDADES
from users.models import Constructora

HOJA_PRECIOS = 'LISTA DE PRECIO'
HOJA_NOMBRES_OBRAS = 'LISTA EMPRESAS REGULARES'
HOJA_NOMBRES_VARIOS = 'LISTA EMPRESAS VARIOS OBRAS'
PRIMERA_COL_OBRA = 4          # columna E (DESCUENTO, si existe, no parece código y se salta)
FILA_ENCABEZADO = 3           # índice 0 -> fila 4 de Excel (la fila 3 es igual pero con celdas "0")
SUFIJO_CATEGORIA = {'11L': 'Laboratorista', '11A': 'Auxiliar', '11B': 'Laboratorista mes'}
_CODIGO_OBRA = re.compile(r'^\d+-\d+$')


@dataclass
class ServicioExcel:
    fila: int
    categoria: str
    codigo: str
    nombre: str
    norma: str
    precio_lista: Decimal | None
    precios: dict = field(default_factory=dict)   # codigo_obra_excel -> Decimal


@dataclass
class LecturaExcel:
    categorias: list            # [(codigo, nombre)]
    servicios: list             # [ServicioExcel]
    obras: list                 # códigos de obra del Excel, en orden de columna
    nombres_obras: dict         # codigo_obra -> nombre proyecto (si la hoja existe)
    nombres_empresas: dict      # codigo_empresa -> razón social
    log: list                   # decisiones de limpieza, en texto


def _norm(texto):
    return re.sub(r'\s+', ' ', str(texto).replace('\xa0', ' ')).strip()


def _codigo(texto):
    return re.sub(r'[,.]', '-', _norm(texto)).replace(' ', '')


def _corregir_prefijo(codigo, categoria):
    """
    Error de digitación frecuente: dentro de la categoría "1" aparece "13-8"
    (falta el guion: es "1-3-8"). Si el primer segmento del código no es el
    número de la categoría pero empieza por él, se reinserta el guion.
    """
    m = re.match(r'\d+', categoria)
    if not m:
        return codigo
    num_cat = m.group()
    primero, sep, resto = codigo.partition('-')
    if primero.isdigit() and primero != num_cat and primero.startswith(num_cat) and len(primero) > len(num_cat):
        return f'{num_cat}-{primero[len(num_cat):]}' + (sep + resto if sep else '')
    return codigo


def _unir_normas(a, b):
    """'NTC 176 / INV E 223' sin repetir partes: une por ' / ' y deduplica."""
    partes = []
    for trozo in f'{a} / {b}'.split('/'):
        trozo = trozo.strip()
        if trozo and trozo not in partes:
            partes.append(trozo)
    return ' / '.join(partes)


def _decimal(valor):
    if isinstance(valor, bool) or not isinstance(valor, (int, float, Decimal)):
        return None
    d = Decimal(str(valor))
    return d if d > 0 else None


def leer_lista_precios(ruta, hoja=HOJA_PRECIOS):
    """Lee el libro y devuelve una LecturaExcel ya limpia."""
    import openpyxl  # import tardío: solo hace falta al importar

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if hoja not in wb.sheetnames:
        raise ValueError(f'El libro no tiene la hoja "{hoja}". Hojas: {wb.sheetnames}')
    filas = list(wb[hoja].iter_rows(values_only=True))
    log = []

    # ── columnas de obra ────────────────────────────────────────────────
    encabezado = filas[FILA_ENCABEZADO]
    respaldo = filas[FILA_ENCABEZADO - 1] if FILA_ENCABEZADO else ()
    columnas = []                      # [(indice_columna, codigo_obra)]
    for j in range(PRIMERA_COL_OBRA, len(encabezado)):
        celda = encabezado[j] if encabezado[j] not in (None, '', 0) else (respaldo[j] if j < len(respaldo) else None)
        if celda in (None, '', 0):
            continue
        cod = _codigo(celda)
        if _CODIGO_OBRA.match(cod):
            columnas.append((j, cod))
        elif cod.upper().strip() not in ('DESCUENTO', 'PRECIO LISTA'):
            log.append(f'columna {j + 1}: el encabezado "{cod}" no es un código de obra (empresa-obra): se omite')
    repetidas = [c for c, n in Counter(c for _, c in columnas).items() if n > 1]
    for c in repetidas:
        log.append(f'columna de obra "{c}" aparece más de una vez: se toma el primer precio no vacío')

    # ── categorías y servicios ──────────────────────────────────────────
    categorias, servicios, cat = [], [], None
    for i, fila in enumerate(filas[FILA_ENCABEZADO + 1:], FILA_ENCABEZADO + 2):
        item, desc = fila[0], fila[1] if len(fila) > 1 else None
        if item is None and desc is None:
            continue
        item = _norm(item) if item is not None else ''
        desc = _norm(desc) if desc else ''
        if not item or not desc:
            continue
        norma = fila[2] if len(fila) > 2 else None
        precio_lista = fila[3] if len(fila) > 3 else None
        if '-' not in item and desc.isupper() and norma is None and precio_lista is None:
            cat = _codigo(item)
            categorias.append((cat, desc))
            continue
        if cat is None:
            log.append(f'fila {i}: servicio "{item}" antes de la primera categoría, omitido')
            continue
        codigo = _corregir_prefijo(_codigo(item), cat)
        if codigo != item.replace(' ', ''):
            log.append(f'fila {i}: código "{item}" -> "{codigo}"')
        s = ServicioExcel(
            fila=i, categoria=cat, codigo=codigo, nombre=desc,
            norma='' if norma in (None, 'N/A') else _norm(norma),
            precio_lista=_decimal(precio_lista),
        )
        for j, cod_obra in columnas:
            valor = _decimal(fila[j]) if j < len(fila) else None
            if valor is not None and cod_obra not in s.precios:
                s.precios[cod_obra] = valor
        servicios.append(s)

    # ── 1) mismo nombre: fusionar dentro de la categoría, sufijo entre categorías
    por_nombre, finales = {}, []
    for s in servicios:
        clave = s.nombre.lower()
        previo = por_nombre.get(clave)
        if previo is not None:
            if previo.categoria == s.categoria:
                previo.norma = _unir_normas(previo.norma, s.norma)
                for cod_obra, valor in s.precios.items():
                    previo.precios.setdefault(cod_obra, valor)
                log.append(f'fila {s.fila}: "{s.codigo}" omitido, mismo nombre que "{previo.codigo}" (normas y precios fusionados)')
                continue
            sufijo = SUFIJO_CATEGORIA.get(s.categoria) or f'cód. {s.codigo}'
            log.append(f'fila {s.fila}: "{s.codigo}" renombrado con sufijo "({sufijo})" por chocar con "{previo.codigo}" de otra categoría')
            s.nombre = f'{s.nombre} ({sufijo})'
        por_nombre[s.nombre.lower()] = s
        finales.append(s)

    # ── 2) mismo código con distinto nombre: sufijo -2, -3, ...
    vistos = Counter()
    for s in finales:
        vistos[s.codigo] += 1
        if vistos[s.codigo] > 1:
            nuevo = f'{s.codigo}-{vistos[s.codigo]}'
            log.append(f'fila {s.fila}: código "{s.codigo}" repetido con otro nombre -> "{nuevo}"')
            s.codigo = nuevo

    for s in finales:
        s.norma = s.norma[:100]
        s.nombre = s.nombre[:255]

    nombres_obras, nombres_empresas = _leer_nombres(wb)
    obras = []
    for _, cod in columnas:
        if cod not in obras:
            obras.append(cod)
    return LecturaExcel(categorias, finales, obras, nombres_obras, nombres_empresas, log)


def _leer_nombres(wb):
    """Hoja LISTA EMPRESAS REGULARES: fila 1 = (cod_empresa, razón social)..., fila 2 = (cod_obra, proyecto)..."""
    if HOJA_NOMBRES_OBRAS not in wb.sheetnames:
        return {}, {}
    filas = list(wb[HOJA_NOMBRES_OBRAS].iter_rows(values_only=True, max_row=2))
    def pares(fila):
        celdas = [c for c in fila if c is not None]
        out = {}
        for k in range(0, len(celdas) - 1, 2):
            cod, nombre = _codigo(celdas[k]), _norm(celdas[k + 1])
            if nombre and not nombre.startswith('Columna'):
                out[cod] = nombre
        return out
    empresas = pares(filas[0]) if filas else {}
    obras = pares(filas[1]) if len(filas) > 1 else {}
    # Clientes varios: una fila por (código, texto); la primera vez es la razón social y la segunda el proyecto
    if HOJA_NOMBRES_VARIOS in wb.sheetnames:
        for fila in wb[HOJA_NOMBRES_VARIOS].iter_rows(values_only=True):
            celdas = [c for c in fila if c not in (None, '')]
            if len(celdas) < 2:
                continue
            cod, texto = _codigo(celdas[0]), _norm(celdas[1])
            if not _CODIGO_OBRA.match(cod) or not texto or texto.startswith('Columna'):
                continue
            if cod not in empresas:
                empresas[cod] = texto
            elif cod not in obras and texto.lower() != empresas[cod].lower():
                obras[cod] = texto
    return obras, empresas


# ═══════════════════════════════════════════════════════════════════════════
# Resolución de obras del Excel -> Obra del sistema
# ═══════════════════════════════════════════════════════════════════════════

def _clave_obra(nombre):
    """
    "IBA 14-1", "IBA14-1", "iba 14 - 1" -> "14-1". Las obras migradas desde
    WordPress se llaman así (el código del Excel con el prefijo de la ciudad).
    """
    s = _norm(nombre).upper().replace(' ', '')
    for prefijo in CIUDADES.values():
        if s.startswith(prefijo):
            s = s[len(prefijo):]
            break
    return s.lstrip('-')


def resolver_obra(ciudad, codigo_excel, lectura, overrides=None, columnas_de_la_empresa=None):
    """
    Devuelve (obra, motivo, candidatas). obra es None si no se pudo resolver.

    Orden de búsqueda:
    1. --obra CODIGO_EXCEL=CODIGO_OBRA pasado por el usuario.
    2. Constructora de la ciudad con código IBA8 / 8 / IBA-8 (prefijo de la
       ciudad + número de empresa) o con la razón social del Excel; dentro de
       ella, la obra cuyo codigo_obra es IBA8-1 / 8-1, o cuyo NOMBRE es el
       código del Excel ("IBA 8-1"), o cuyo nombre es el proyecto del Excel.
    3. Si la constructora tiene una sola obra y el Excel también tiene una
       sola columna para esa empresa, se asume que es esa (motivo "supuesto").
    """
    overrides = overrides or {}
    if codigo_excel in overrides:
        obra = Obra.objects.filter(
            codigo_obra__iexact=overrides[codigo_excel], constructora__ciudad__iexact=ciudad,
        ).select_related('constructora').first()
        return obra, ('--obra' if obra else f'--obra: no existe "{overrides[codigo_excel]}" en {ciudad}'), []

    emp, _, num = codigo_excel.partition('-')
    prefijo = CIUDADES.get(ciudad, '')
    if emp == '0':
        return _resolver_varios(ciudad, codigo_excel, lectura, prefijo)
    codigos_emp = {f'{prefijo}{emp}', emp, f'{prefijo}-{emp}', f'{prefijo}{emp.zfill(2)}'}
    razon = lectura.nombres_empresas.get(emp)
    constructoras = _constructoras_de(ciudad, codigos_emp, razon)
    if not constructoras:
        return None, f'ninguna constructora en {ciudad} con código {", ".join(sorted(codigos_emp))}' + (f' ni razón social "{razon}"' if razon else ''), []

    proyecto = lectura.nombres_obras.get(codigo_excel)
    candidatas = []
    for c in constructoras:
        obras = list(Obra.objects.filter(constructora=c).order_by('codigo_obra'))
        candidatas.extend(obras)
        codigos_obra = {x.upper() for x in (f'{c.codigo}-{num}', codigo_excel, f'{prefijo}{codigo_excel}')}
        for o in obras:
            if o.codigo_obra.upper() in codigos_obra:
                return o, f'código de obra {o.codigo_obra}' + (f' (en {c.codigo} {c.nombre[:30]})' if len(constructoras) > 1 else ''), candidatas
        for o in obras:
            if _clave_obra(o.nombre) == codigo_excel:
                return o, f'nombre de obra "{o.nombre}"', candidatas
        if proyecto:
            for o in obras:
                if o.nombre.strip().lower() == proyecto.lower():
                    return o, f'nombre de proyecto "{proyecto}"', candidatas
    if len(candidatas) == 1 and (columnas_de_la_empresa or 1) == 1:
        return candidatas[0], f'única obra de {candidatas[0].constructora.nombre} y única columna de la empresa en el Excel (supuesto)', candidatas
    n_cols = columnas_de_la_empresa or 1
    if len(candidatas) == 0:
        return None, f'{constructoras[0].nombre} no tiene obras en el sistema (use --crear-obras o --obra {codigo_excel}=CODIGO_OBRA)', candidatas
    return None, (
        f'{constructoras[0].nombre} tiene {len(candidatas)} obra(s) en el sistema y el Excel {n_cols} columna(s); '
        f'ninguna se llama "{prefijo} {codigo_excel}" (use --crear-obras o --obra {codigo_excel}=CODIGO_OBRA)'
    ), candidatas


def _resolver_varios(ciudad, codigo_excel, lectura, prefijo):
    """
    "0-10" es el cliente varios número 10: en el portal su constructora y su obra
    se llaman BOG0-10 (a veces la constructora lleva sufijo B, o la obra cuelga
    de otra constructora pero conserva el código). Si hay varias obras con el
    código (WordPress lo reutilizó), desempata la razón social del Excel.
    """
    cod = f'{prefijo}{codigo_excel}'
    razon = lectura.nombres_empresas.get(codigo_excel)
    obras = list(
        Obra.objects.filter(constructora__ciudad__iexact=ciudad).filter(
            Q(codigo_obra__iexact=cod)
            | Q(codigo_obra__iregex=rf'^{re.escape(cod)}-\d+$')
            | Q(constructora__codigo__iregex=rf'^{re.escape(cod)}[A-Z]?$')
        ).select_related('constructora').distinct().order_by('codigo_obra')
    )
    excel = f' (el Excel lo llama "{razon}")' if razon else ' (el Excel no trae su nombre)'
    if not obras:
        ajena = Constructora.objects.filter(codigo__iexact=cod).exclude(ciudad__iexact=ciudad).first()
        if ajena:
            return None, (f'el código {cod} ya lo tiene "{ajena.nombre}" en {ajena.ciudad or "sin ciudad"}{excel}: '
                          f'corrige la ciudad de esa empresa o use --obra {codigo_excel}=CODIGO_OBRA'), []
        if not razon:
            return None, f'ningún cliente varios {cod} en {ciudad}{excel}: sin nombre no se puede crear; use --obra {codigo_excel}=CODIGO_OBRA', []
        return None, f'ningún cliente varios {cod} en {ciudad}{excel}; use --crear-obras (crea empresa y obra) o --obra {codigo_excel}=CODIGO_OBRA', []
    if len(obras) == 1:
        o = obras[0]
        aviso = ''
        if razon and not (misma_empresa(razon, o.constructora.nombre) or misma_empresa(razon, o.nombre)):
            aviso = f' ⚠ el Excel dice "{razon}" y en el portal es "{o.constructora.nombre}": confirmar'
        return o, f'cliente varios {cod}{aviso}', obras
    if razon:
        afines = [o for o in obras if misma_empresa(razon, o.constructora.nombre) or misma_empresa(razon, o.nombre)]
        if len(afines) == 1:
            return afines[0], f'cliente varios {cod}, razón social "{razon}"', obras
    return None, (
        f'{len(obras)} obras con el código {cod} en {ciudad}'
        + (f' y la razón social "{razon}" no desempata' if razon else excel)
        + f' (use --obra {codigo_excel}=CODIGO_OBRA)'
    ), obras


def _constructoras_de(ciudad, codigos_emp, razon):
    """
    El CÓDIGO manda: si hay constructora con ese código (IBA14), es esa aunque
    su razón social no coincida con el Excel (en WordPress hay razones sociales
    mal capturadas). La razón social solo se usa cuando ningún código coincide.
    Los sufijos B, C (códigos reutilizados) también cuentan como el mismo código.
    """
    base = Q(ciudad__iexact=ciudad)
    # IBA10 y sus sufijos IBA10B, IBA10C (código reutilizado en WordPress para otra
    # empresa) se consideran juntos: la obra "IBA10-1" puede estar en cualquiera.
    patron = r'^(' + '|'.join(sorted(codigos_emp)) + r')[A-Z]?$'
    por_codigo = list(Constructora.objects.filter(base, codigo__regex=patron).order_by('codigo'))
    if por_codigo:
        return por_codigo
    if razon:
        return list(Constructora.objects.filter(base, nombre__iexact=razon))
    return []


def constructora_de(ciudad, codigo_excel, lectura):
    """Constructora a la que pertenece una columna del Excel (o None). Misma regla que resolver_obra."""
    emp = codigo_excel.partition('-')[0]
    prefijo = CIUDADES.get(ciudad, '')
    if emp == '0':
        cod = f'{prefijo}{codigo_excel}'
        return Constructora.objects.filter(
            ciudad__iexact=ciudad, codigo__iregex=rf'^{re.escape(cod)}[A-Z]?$'
        ).order_by('codigo').first()
    codigos_emp = {f'{prefijo}{emp}', emp, f'{prefijo}-{emp}', f'{prefijo}{emp.zfill(2)}'}
    lista = _constructoras_de(ciudad, codigos_emp, lectura.nombres_empresas.get(emp))
    return lista[0] if lista else None


def crear_obra(ciudad, codigo_excel, lectura):
    """
    Crea la obra de una columna del Excel bajo su constructora, con la misma
    convención de las obras migradas: nombre "IBA 8-1" (o el proyecto del
    Excel si se conoce) y codigo_obra "IBA8-1". Devuelve None si no hay constructora.
    """
    c = constructora_de(ciudad, codigo_excel, lectura)
    prefijo = CIUDADES.get(ciudad, '')
    if codigo_excel.startswith('0-'):
        # Cliente varios nuevo: empresa + obra con el código del Excel, solo si el libro trae su nombre
        razon = lectura.nombres_empresas.get(codigo_excel)
        if c is None:
            # Sin nombre no se inventa; y el código es único en todo el portal (otra ciudad lo bloquea)
            if not razon or Constructora.objects.filter(codigo__iexact=f'{prefijo}{codigo_excel}').exists():
                return None
            c = Constructora.objects.create(codigo=f'{prefijo}{codigo_excel}', nombre=razon, ciudad=ciudad)
        proyecto = lectura.nombres_obras.get(codigo_excel) or razon or c.nombre
        codigo_obra = f'{prefijo}{codigo_excel}'
        if Obra.objects.filter(codigo_obra__iexact=codigo_obra).exists():
            codigo_obra = f'{codigo_obra}-1'
        return Obra.objects.create(constructora=c, nombre=proyecto, codigo_obra=codigo_obra)
    if c is None:
        return None
    num = codigo_excel.partition('-')[2]
    proyecto = lectura.nombres_obras.get(codigo_excel)
    nombre = f'{proyecto} ({prefijo} {codigo_excel})' if proyecto else f'{prefijo} {codigo_excel}'.strip()
    codigo_obra = f'{c.codigo}-{num}'
    if Obra.objects.filter(codigo_obra__iexact=codigo_obra).exists():
        codigo_obra = f'{c.codigo}-{codigo_excel}'
    return Obra.objects.create(constructora=c, nombre=nombre, codigo_obra=codigo_obra)
