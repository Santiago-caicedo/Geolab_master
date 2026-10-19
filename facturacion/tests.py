"""
Tests del módulo de facturación.

    python manage.py test facturacion
"""

import json
from decimal import Decimal
from unittest import mock

from django.test import SimpleTestCase, TestCase

from .models import CategoriaServicio, TipoServicio
from .orden import clave_orden


class ClaveOrdenTests(SimpleTestCase):
    """La clave debe ordenar los códigos como números, no como texto."""

    def test_numeros_simples(self):
        codigos = ['10', '2', '1', '14', '9']
        self.assertEqual(
            sorted(codigos, key=clave_orden), ['1', '2', '9', '10', '14'],
        )

    def test_sufijos_de_letras(self):
        codigos = ['9', '8V', '8A', '10', '11L', '11A', '11B', '12']
        self.assertEqual(
            sorted(codigos, key=clave_orden),
            ['8A', '8V', '9', '10', '11A', '11B', '11L', '12'],
        )

    def test_codigos_compuestos(self):
        codigos = ['1-10', '1-2', '1-1', '1-3-5-2', '1-3-5', '1-1-1', '1-3']
        self.assertEqual(
            sorted(codigos, key=clave_orden),
            ['1-1', '1-1-1', '1-2', '1-3', '1-3-5', '1-3-5-2', '1-10'],
        )

    def test_letras_pegadas_al_numero(self):
        self.assertLess(clave_orden('11-10L'), clave_orden('11-10LV'))
        self.assertLess(clave_orden('8-1A'), clave_orden('8-10A'))

    def test_codigo_corto_antes_que_sus_derivados(self):
        codigos = ['11-1-1L', '11-1L', '11-2L', '11-10L', '11-1-7L']
        self.assertEqual(
            sorted(codigos, key=clave_orden),
            ['11-1L', '11-1-1L', '11-1-7L', '11-2L', '11-10L'],
        )

    def test_codigo_vacio_no_rompe(self):
        self.assertEqual(clave_orden(''), '')
        self.assertEqual(clave_orden(None), '')


class OrdenEnBaseDeDatosTests(TestCase):
    """Meta.ordering usa clave_orden, que se calcula sola en save()."""

    def test_categorias_en_orden_natural(self):
        for codigo in ['10', '2', '11L', '1', '8V', '8A', '9', '14']:
            CategoriaServicio.objects.create(codigo=codigo, nombre=f'CAT {codigo}')
        self.assertEqual(
            list(CategoriaServicio.objects.values_list('codigo', flat=True)),
            ['1', '2', '8A', '8V', '9', '10', '11L', '14'],
        )

    def test_servicios_por_categoria_y_codigo(self):
        c1 = CategoriaServicio.objects.create(codigo='1', nombre='CONCRETOS')
        c10 = CategoriaServicio.objects.create(codigo='10', nombre='PERFORACIONES')
        TipoServicio.objects.create(categoria=c10, codigo='10-1', nombre='s10-1')
        TipoServicio.objects.create(categoria=c1, codigo='1-10', nombre='s1-10')
        TipoServicio.objects.create(categoria=c1, codigo='1-2', nombre='s1-2')
        TipoServicio.objects.create(categoria=c1, codigo='1-1', nombre='s1-1')
        self.assertEqual(
            list(TipoServicio.objects.values_list('codigo', flat=True)),
            ['1-1', '1-2', '1-10', '10-1'],
        )

    def test_clave_se_recalcula_al_cambiar_codigo(self):
        cat = CategoriaServicio.objects.create(codigo='3', nombre='X')
        cat.codigo = '12'
        cat.save()
        cat.refresh_from_db()
        self.assertEqual(cat.clave_orden, '0012')


from django.urls import reverse

from core.models import Obra
from users.models import Constructora, FuncionarioGeolab, UsuarioBase

from .models import Factura, Impuesto, PrecioServicio, RegistroServicio
from .sede import SESSION_KEY


class CiudadFacturacionTests(TestCase):
    """El módulo exige elegir ciudad y filtra todo por ella."""

    @classmethod
    def setUpTestData(cls):
        cls.staff = UsuarioBase.objects.create_user(
            'admin_fact', password='clave-de-prueba', es_geolab=True,
        )
        FuncionarioGeolab.objects.create(user=cls.staff, area='admin')

        cls.c_buc = Constructora.objects.create(nombre='Buc SAS', codigo='B1', ciudad='Bucaramanga')
        cls.c_bog = Constructora.objects.create(nombre='Bog SAS', codigo='G1', ciudad='Bogotá')
        cls.c_sin = Constructora.objects.create(nombre='Sin ciudad', codigo='S1', ciudad='')
        cls.o_buc = Obra.objects.create(nombre='Obra Buc', codigo_obra='B1-1', constructora=cls.c_buc)
        cls.o_bog = Obra.objects.create(nombre='Obra Bog', codigo_obra='G1-1', constructora=cls.c_bog)

        cat = CategoriaServicio.objects.create(codigo='1', nombre='CONCRETOS')
        cls.srv = TipoServicio.objects.create(categoria=cat, codigo='1-1', nombre='Compresión')
        for obra in (cls.o_buc, cls.o_bog):
            RegistroServicio.objects.create(
                obra=obra, tipo_servicio=cls.srv, fecha_realizacion='2026-09-01',
                cantidad=1, precio_unitario_congelado=1000, creado_por=cls.staff,
            )

    def setUp(self):
        self.client.login(username='admin_fact', password='clave-de-prueba')

    def _elegir(self, ciudad):
        session = self.client.session
        session[SESSION_KEY] = ciudad
        session.save()

    def test_sin_ciudad_redirige_al_selector(self):
        r = self.client.get(reverse('dashboard_facturacion'))
        self.assertEqual(r.status_code, 302)
        self.assertIn(reverse('seleccionar_ciudad_facturacion'), r.url)
        self.assertIn('next=', r.url)

    def test_api_sin_ciudad_devuelve_403_json(self):
        r = self.client.get(reverse('api_get_obras_facturacion'), {'constructora_id': self.c_buc.pk})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(r.json()['status'], 'error')

    def test_selector_lista_ciudades_y_avisa_sin_ciudad(self):
        r = self.client.get(reverse('seleccionar_ciudad_facturacion'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual([c for c, _ in r.context['ciudades']], ['Bogotá', 'Bucaramanga'])
        self.assertEqual(r.context['sin_ciudad'], 1)

    def test_elegir_ciudad_guarda_en_sesion_y_respeta_next(self):
        destino = reverse('historico_facturacion')
        r = self.client.post(
            reverse('seleccionar_ciudad_facturacion'),
            {'ciudad': 'Bucaramanga', 'next': destino},
        )
        self.assertRedirects(r, destino)
        self.assertEqual(self.client.session[SESSION_KEY], 'Bucaramanga')

    def test_ciudad_invalida_no_se_acepta(self):
        r = self.client.post(reverse('seleccionar_ciudad_facturacion'), {'ciudad': 'Marte'})
        self.assertEqual(r.status_code, 200)
        self.assertNotIn(SESSION_KEY, self.client.session)

    def test_next_externo_se_ignora(self):
        r = self.client.post(
            reverse('seleccionar_ciudad_facturacion'),
            {'ciudad': 'Bogotá', 'next': 'https://evil.example/x'},
        )
        self.assertRedirects(r, reverse('dashboard_facturacion'), fetch_redirect_response=False)

    def test_dashboard_cuenta_solo_la_ciudad(self):
        self._elegir('Bucaramanga')
        r = self.client.get(reverse('dashboard_facturacion'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['total_constructoras'], 1)
        self.assertEqual(r.context['total_obras'], 1)
        self.assertEqual(r.context['por_facturar_n'], 1)


    def test_vista_previa_entrega_numeros_validos_al_javascript(self):
        # Con es-co Django escribe "1000,00": en JavaScript eso rompía el script y el IVA no se calculaba
        from .models import Impuesto
        Impuesto.objects.create(nombre='IVA', porcentaje='19.00')
        self._elegir('Bucaramanga')
        r = self.client.get(reverse('generar_factura'),
                            {'obra': self.o_buc.pk, 'fecha_inicio': '2026-09-01', 'fecha_fin': '2026-09-30'})
        self.assertContains(r, 'const subtotalNormales = Number("1000.00");')
        self.assertContains(r, 'const ivaPct = Number("19.00");')
        self.assertNotContains(r, '1000,00;')

    def test_panel_clientes_sin_anuladas_y_doce_meses(self):
        from datetime import timedelta
        from django.utils import timezone
        hoy = timezone.localdate()
        base = dict(constructora=self.c_buc, obra=self.o_buc, fecha_inicio_periodo=hoy, fecha_fin_periodo=hoy,
                    fecha_emision=hoy, subtotal=100, monto_iva=0)
        Factura.objects.create(total=100, estado='PAGADA', **base)
        Factura.objects.create(total=500, estado='ANULADA', **base)   # antes sacaba al cliente del ranking
        Factura.objects.create(total=40, estado='PENDIENTE', **base)
        self._elegir('Bucaramanga')
        ctx = self.client.get(reverse('dashboard_facturacion')).context
        self.assertEqual([(c['constructora__nombre'], c['total']) for c in ctx['top_clientes']], [('Buc SAS', 140)])
        self.assertEqual((ctx['total_mes'], ctx['facturas_mes']), (140, 2))
        self.assertEqual((ctx['por_cobrar_total'], ctx['por_cobrar_n']), (40, 1))
        self.assertEqual(len(ctx['serie']), 12)
        self.assertTrue(ctx['serie'][-1]['actual'])
        self.assertEqual(ctx['serie'][-1]['total'], 140)
        self.assertEqual(ctx['obras_por_facturar'][0]['n'], 1)

    def test_historico_y_formularios_filtran(self):
        self._elegir('Bucaramanga')
        r = self.client.get(reverse('historico_facturacion'))
        self.assertEqual(r.context['total_registros'], 1)
        self.assertEqual(list(r.context['form'].fields['constructora'].queryset), [self.c_buc])
        r = self.client.get(reverse('generar_factura'))
        self.assertEqual(list(r.context['form'].fields['constructora'].queryset), [self.c_buc])

    def test_api_obras_no_devuelve_otra_ciudad(self):
        self._elegir('Bucaramanga')
        r = self.client.get(reverse('api_get_obras_facturacion'), {'constructora_id': self.c_bog.pk})
        self.assertEqual(r.json()['obras'], [])
        r = self.client.get(reverse('api_get_obras_facturacion'), {'constructora_id': self.c_buc.pk})
        self.assertEqual([o['id'] for o in r.json()['obras']], [self.o_buc.pk])

    def test_objetos_de_otra_ciudad_dan_404(self):
        self._elegir('Bucaramanga')
        self.assertEqual(
            self.client.get(reverse('gestionar_precios_obra', args=[self.o_bog.pk])).status_code, 404,
        )
        self.assertEqual(
            self.client.get(reverse('repositorio_obras_cliente', args=[self.c_bog.pk])).status_code, 404,
        )
        reg_bog = RegistroServicio.objects.get(obra=self.o_bog)
        self.assertEqual(
            self.client.get(reverse('editar_registro_facturacion', args=[reg_bog.pk])).status_code, 404,
        )

    def test_crear_registro_rechaza_obra_de_otra_ciudad(self):
        self._elegir('Bucaramanga')
        r = self.client.post(
            reverse('crear_registro_facturacion'),
            data='{"obra_id": %d, "fecha_realizacion": "2026-09-01", "servicios": [{"tipo_id": %d, "cantidad": 1}]}'
                 % (self.o_bog.pk, self.srv.pk),
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 404)

    def test_sidebar_muestra_ciudad_activa(self):
        self._elegir('Bucaramanga')
        r = self.client.get(reverse('dashboard_facturacion'))
        self.assertContains(r, 'Bucaramanga')
        self.assertContains(r, reverse('seleccionar_ciudad_facturacion'))


import tempfile
from decimal import Decimal
from pathlib import Path

from django.core.management import call_command
from django.core.management.base import CommandError
from io import StringIO

from .importar_excel import leer_lista_precios, resolver_obra


def _libro_de_prueba(ruta):
    """Libro mínimo con la estructura de BASE DATOS <CIUDAD>.xlsm."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'LISTA DE PRECIO'
    ws.append(['', '', 'LISTA DE PRECIOS 2023'])
    ws.append([])
    ws.append(['ÍTEM', 'DESCRIPCIÓN', 'NORMA', 'PRECIO LISTA', 'DESCUENTO ', '8-1', 0, '10-2', '8-1'])
    ws.append(['ÍTEM', 'DESCRIPCIÓN', 'NORMA', 'PRECIO LISTA', 'DESCUENTO ', '8-1', None, '10-2', '8-1'])
    ws.append(['\xa0       1', 'CONCRETOS'])
    ws.append(['1-1', 'Compresión de cilindros', 'NTC 673', 4500, -4500, 5000, None, 5500, None])
    ws.append(['1-2', 'Diseño de mezclas', 'N/A', 420000, -420000, None, None, 0, 350000])
    ws.append(['1-2', 'Diseño de mezclas', 'NTC 999', None, None, 111, None, None, None])   # dup nombre: fusiona
    ws.append(['13-8', 'Tracción indirecta', 'NTC 722', 20000, None, 21000, None, None, None])  # falta guion
    ws.append(['2', 'SUELOS'])
    ws.append(['2.1', 'Humedad', 'INV E-122', 5000, None, 6000, None, 6500, None])
    ws.append(['2-2', 'Compresión de cilindros', 'X', 1, None, None, None, None, None])       # mismo nombre otra cat.
    ws2 = wb.create_sheet('LISTA EMPRESAS REGULARES')
    ws2.append(['8', 'DYCO SAS', '10', 'ASFALTEMOS', '11', 'Columna1'])
    ws2.append(['8-1', 'AMBALA', '10-2', 'JORDAN'])
    wb.save(ruta)


def _libro_bogota(ruta):
    """Libro como BASE DATOS BOGOTA.xlsm: sin columna DESCUENTO (las obras empiezan en E),
    un encabezado suelto "1" y clientes varios 0-N con su hoja de nombres."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'LISTA DE PRECIO'
    ws.append(['', '', 'LISTA DE PRECIOS 2019'])
    ws.append([])
    ws.append(['ÍTEM', 'DESCRIPCIÓN', 'NORMA', 'PRECIO LISTA', '1-1', '0-10', '1', '0-2', '0-22', '0-34'])
    ws.append(['ÍTEM', 'DESCRIPCIÓN', 'NORMA', 'PRECIO LISTA', '1-1', '0-10', '1', '0-2', '0-22', '0-34'])
    ws.append(['\xa0       1', 'CONCRETOS'])
    ws.append(['1-1', 'Compresión de cilindros', 'NTC 673', 5500, 4500, 4800, 9, 5000, 5100, 5200])
    ws2 = wb.create_sheet('LISTA EMPRESAS REGULARES')
    ws2.append(['1', 'MAB INGENIERIA DE VALORES'])
    ws2.append(['1-1', 'CONTRATO IDU 1286'])
    ws3 = wb.create_sheet('LISTA EMPRESAS VARIOS OBRAS')
    for fila in (['0-10', 'CONSORCIO DC 2025'], ['0-10', 'MALLA VIAL LOCAL'], ['0-2', 'CONSORCIO INFRACENTRO DP'],
                 ['0-2', 'CONSORCIO INFRACENTRO DP'], ['0-22', 'Columna1'], ['0-34', 'CONSORCIO ESTACIONES NAVARRA'],
                 ['0-34', 'RENOVACION ESTACION NAVARRA']):
        ws3.append(fila)
    wb.save(ruta)


class ImportarBogotaTests(TestCase):
    """Bogotá: columnas de obra desde la E y clientes varios 0-N (constructora/obra BOG0-N)."""

    @classmethod
    def setUpTestData(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = str(Path(cls.tmp.name) / 'bogota.xlsx')
        _libro_bogota(cls.ruta)
        cls.lectura = leer_lista_precios(cls.ruta)
        mab = Constructora.objects.create(codigo='BOG1', nombre='MAB INGENIERIA DE SERVICIOS SA', ciudad='Bogotá')
        cls.o11 = Obra.objects.create(constructora=mab, nombre='CONTRATO IDU No.1286', codigo_obra='BOG1-1')
        # 0-10: la obra BOG0-10 cuelga de una constructora con otro código (así vino de WordPress)
        dc = Constructora.objects.create(codigo='BOG0-0B', nombre='CONSORCIO DC 2025', ciudad='Bogotá')
        cls.o010 = Obra.objects.create(constructora=dc, nombre='0-10 EJECUTAR A MONTO', codigo_obra='BOG0-10')
        # 0-2: tres obras con el mismo código; solo una es de INFRACENTRO
        proar = Constructora.objects.create(codigo='BOG0-2', nombre='CONSORCIO OBRAS ACUEDUCTO PROAR', ciudad='Bogotá')
        Obra.objects.create(constructora=proar, nombre='ALCANTARILLADO ZONA 1', codigo_obra='BOG0-2')
        infra = Constructora.objects.create(codigo='BOG0-2B', nombre='CONSORCIO INFRACENTRO DP', ciudad='Bogotá')
        cls.o02 = Obra.objects.create(constructora=infra, nombre='INTERVENTORIA IDU 1523-2025', codigo_obra='BOG0-2')
        # 0-34 existe pero en otra ciudad: el código es único en el portal
        Constructora.objects.create(codigo='BOG0-34', nombre='OTRA', ciudad='Ibagué')

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        super().tearDownClass()

    def test_lee_columnas_desde_la_e_y_nombres_de_varios(self):
        self.assertEqual(self.lectura.obras, ['1-1', '0-10', '0-2', '0-22', '0-34'])
        self.assertTrue(any('"1" no es un código de obra' in l for l in self.lectura.log))
        self.assertEqual(self.lectura.nombres_empresas['0-10'], 'CONSORCIO DC 2025')
        self.assertEqual(self.lectura.nombres_obras['0-10'], 'MALLA VIAL LOCAL')
        self.assertNotIn('0-22', self.lectura.nombres_empresas)
        self.assertEqual({s.codigo: s.precios['1-1'] for s in self.lectura.servicios}, {'1-1': Decimal('4500')})

    def test_resuelve_clientes_varios(self):
        self.assertEqual(resolver_obra('Bogotá', '1-1', self.lectura)[0], self.o11)
        obra, motivo, _ = resolver_obra('Bogotá', '0-10', self.lectura)
        self.assertEqual((obra, 'cliente varios BOG0-10' in motivo, '⚠' in motivo), (self.o010, True, False))
        obra, motivo, cand = resolver_obra('Bogotá', '0-2', self.lectura)
        self.assertEqual(obra, self.o02)
        self.assertIn('razón social', motivo)
        obra, motivo, _ = resolver_obra('Bogotá', '0-22', self.lectura)
        self.assertIsNone(obra)
        self.assertIn('sin nombre', motivo)
        obra, motivo, _ = resolver_obra('Bogotá', '0-34', self.lectura)
        self.assertIsNone(obra)
        self.assertIn('Ibagué', motivo)

    def test_crear_obras_crea_cliente_varios_con_nombre(self):
        out = StringIO()
        call_command('importar_catalogo_excel', self.ruta, '--ciudad', 'Bogotá', '--crear-obras', stdout=out)
        salida = out.getvalue()
        self.assertIn('✓    1-1', salida)
        self.assertIn('✓   0-10', salida)
        self.assertIn('✓    0-2', salida)
        self.assertIn('✗   0-22', salida)      # sin nombre: no se inventa
        self.assertIn('✗   0-34', salida)      # el código es de otra ciudad: no se crea
        self.assertFalse(Constructora.objects.filter(codigo='BOG0-22').exists())
        self.assertEqual(Constructora.objects.filter(codigo='BOG0-34').count(), 1)
        self.assertEqual(PrecioServicio.objects.filter(obra=self.o02, precio=5000).count(), 1)
        self.assertEqual(PrecioServicio.objects.filter(obra=self.o010, precio=4800).count(), 1)
        # misma corrida otra vez: idempotente
        out = StringIO()
        call_command('importar_catalogo_excel', self.ruta, '--ciudad', 'Bogotá', '--crear-obras', stdout=out)
        self.assertIn('0 nuevos', out.getvalue())


class LecturaExcelTests(SimpleTestCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = str(Path(cls.tmp.name) / 'base.xlsx')
        _libro_de_prueba(cls.ruta)
        cls.lectura = leer_lista_precios(cls.ruta)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        super().tearDownClass()

    def test_categorias_y_columnas(self):
        self.assertEqual(self.lectura.categorias, [('1', 'CONCRETOS'), ('2', 'SUELOS')])
        self.assertEqual(self.lectura.obras, ['8-1', '10-2'])          # la "8-1" repetida no se duplica
        self.assertEqual(self.lectura.nombres_obras, {'8-1': 'AMBALA', '10-2': 'JORDAN'})
        self.assertEqual(self.lectura.nombres_empresas, {'8': 'DYCO SAS', '10': 'ASFALTEMOS'})

    def test_limpieza_de_codigos_y_nombres(self):
        por_codigo = {s.codigo: s for s in self.lectura.servicios}
        self.assertEqual(set(por_codigo), {'1-1', '1-2', '1-3-8', '2-1', '2-2'})
        self.assertEqual(por_codigo['1-2'].norma, 'NTC 999')                       # 'N/A' + fusión
        self.assertEqual(por_codigo['2-2'].nombre, 'Compresión de cilindros (cód. 2-2)')

    def test_precios_por_obra(self):
        por_codigo = {s.codigo: s for s in self.lectura.servicios}
        self.assertEqual(por_codigo['1-1'].precios, {'8-1': Decimal('5000'), '10-2': Decimal('5500')})
        # 0 se ignora; la columna 8-1 repetida aporta 350000; la fila fusionada aporta 111 solo si faltaba
        self.assertEqual(por_codigo['1-2'].precios, {'8-1': Decimal('350000')})
        self.assertEqual(por_codigo['2-1'].precios, {'8-1': Decimal('6000'), '10-2': Decimal('6500')})

    def test_hoja_inexistente(self):
        with self.assertRaises(ValueError):
            leer_lista_precios(self.ruta, hoja='NO EXISTE')


class ImportarCatalogoExcelTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = str(Path(cls.tmp.name) / 'base.xlsx')
        _libro_de_prueba(cls.ruta)
        dyco = Constructora.objects.create(codigo='IBA8', nombre='DYCO SAS', ciudad='Ibagué')
        cls.ambala = Obra.objects.create(constructora=dyco, nombre='AMBALA', codigo_obra='IBA8-1')
        asf = Constructora.objects.create(codigo='IBA10', nombre='ASFALTEMOS', ciudad='Ibagué')
        cls.jordan = Obra.objects.create(constructora=asf, nombre='Jordan', codigo_obra='IBA10-555')
        Obra.objects.create(constructora=asf, nombre='Otra', codigo_obra='IBA10-556')
        # mismo código en otra ciudad: no debe mezclarse
        cat_buc = CategoriaServicio.objects.create(ciudad='Bucaramanga', codigo='1', nombre='CONCRETOS')
        TipoServicio.objects.create(categoria=cat_buc, codigo='1-1', nombre='Otro servicio con código 1-1')

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        super().tearDownClass()

    def _run(self, *extra):
        out = StringIO()
        call_command('importar_catalogo_excel', self.ruta, '--ciudad', 'ibague', *extra, stdout=out)
        return out.getvalue()

    def test_importa_catalogo_y_precios_resolviendo_obras(self):
        salida = self._run()
        self.assertEqual(TipoServicio.objects.filter(ciudad='Ibagué').count(), 5)
        self.assertEqual(TipoServicio.objects.filter(ciudad='Bucaramanga', codigo='1-1').count(), 1)
        precio = PrecioServicio.objects.get(obra=self.ambala, tipo_servicio__codigo='1-1', tipo_servicio__ciudad='Ibagué')
        self.assertEqual(precio.precio, Decimal('5000'))
        # 10-2 se resolvió por nombre de proyecto (JORDAN) aunque el código no coincide
        self.assertTrue(PrecioServicio.objects.filter(obra=self.jordan, precio=Decimal('5500')).exists())
        self.assertIn('Importación guardada', salida)

    def test_es_idempotente_y_dry_run_no_escribe(self):
        self._run()
        antes = PrecioServicio.objects.count()
        salida = self._run()
        self.assertIn('0 nuevos', salida)
        self.assertEqual(PrecioServicio.objects.count(), antes)
        PrecioServicio.objects.all().delete()
        salida = self._run('--dry-run')
        self.assertIn('DRY-RUN', salida)
        self.assertEqual(PrecioServicio.objects.count(), 0)

    def test_no_sobrescribir_conserva_precio_manual(self):
        self._run()
        ps = PrecioServicio.objects.get(obra=self.ambala, tipo_servicio__codigo='1-1', tipo_servicio__ciudad='Ibagué')
        ps.precio = Decimal('1')
        ps.save()
        self._run('--no-sobrescribir')
        ps.refresh_from_db()
        self.assertEqual(ps.precio, Decimal('1'))
        self._run()
        ps.refresh_from_db()
        self.assertEqual(ps.precio, Decimal('5000'))

    def test_obra_manual_y_no_resuelta(self):
        # ya no coincide ni por código ni por nombre de proyecto: cae en 'única obra'
        Obra.objects.filter(codigo_obra='IBA8-1').update(codigo_obra='IBA8-999', nombre='Bodega')
        salida = self._run()
        self.assertIn('única obra de DYCO SAS', salida)
        salida = self._run('--obra', '8-1=IBA10-556')
        self.assertIn('--obra', salida)
        self.assertTrue(PrecioServicio.objects.filter(obra__codigo_obra='IBA10-556', precio=Decimal('5000')).exists())

    def test_conflicto_de_nombre_aborta(self):
        cat = CategoriaServicio.objects.create(ciudad='Ibagué', codigo='9', nombre='OTRA')
        TipoServicio.objects.create(categoria=cat, codigo='9-1', nombre='Humedad')
        with self.assertRaises(CommandError):
            self._run()
        self.assertEqual(TipoServicio.objects.filter(ciudad='Ibagué').count(), 1)


class CatalogoPorCiudadTests(TestCase):

    def test_mismo_codigo_en_dos_ciudades(self):
        for ciudad in ('Bucaramanga', 'Ibagué'):
            cat = CategoriaServicio.objects.create(ciudad=ciudad, codigo='1', nombre='CONCRETOS')
            TipoServicio.objects.create(categoria=cat, codigo='1-8', nombre=f'Servicio 1-8 de {ciudad}')
        self.assertEqual(TipoServicio.objects.filter(codigo='1-8').count(), 2)
        self.assertEqual(TipoServicio.objects.get(codigo='1-8', ciudad='Ibagué').nombre, 'Servicio 1-8 de Ibagué')

    def test_servicio_hereda_ciudad_de_su_categoria(self):
        cat = CategoriaServicio.objects.create(ciudad='Ibagué', codigo='2', nombre='SUELOS')
        s = TipoServicio.objects.create(categoria=cat, codigo='2-1', nombre='Humedad')
        self.assertEqual(s.ciudad, 'Ibagué')

    def test_formulario_de_servicio_rechaza_duplicado_solo_en_su_ciudad(self):
        from .forms import TipoServicioForm
        cat_i = CategoriaServicio.objects.create(ciudad='Ibagué', codigo='1', nombre='CONCRETOS')
        cat_b = CategoriaServicio.objects.create(ciudad='Bucaramanga', codigo='1', nombre='CONCRETOS')
        TipoServicio.objects.create(categoria=cat_b, codigo='1-1', nombre='Compresión')
        datos = {'categoria': cat_i.pk, 'codigo': '1-1', 'nombre': 'Compresión', 'norma': ''}
        self.assertTrue(TipoServicioForm(datos, ciudad='Ibagué').is_valid())
        form = TipoServicioForm({**datos, 'categoria': cat_b.pk}, ciudad='Bucaramanga')
        self.assertFalse(form.is_valid())
        self.assertIn('codigo', form.errors)
        # el select de categoría solo ofrece las de la ciudad
        self.assertEqual(list(TipoServicioForm(ciudad='Ibagué').fields['categoria'].queryset), [cat_i])


class ResolverObraTests(TestCase):
    """Cruce de columnas del Excel con obras del sistema (nombres tipo 'IBA 8-1')."""

    @classmethod
    def setUpTestData(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = str(Path(cls.tmp.name) / 'base.xlsx')
        _libro_de_prueba(cls.ruta)
        cls.lectura = leer_lista_precios(cls.ruta)
        cls.sika = Constructora.objects.create(codigo='IBA8', nombre='SIKA', ciudad='Ibagué')
        cls.o81 = Obra.objects.create(constructora=cls.sika, nombre='IBA 8-1', codigo_obra='IBA8-5923')
        cls.asf = Constructora.objects.create(codigo='IBA10', nombre='ASFALTEMOS', ciudad='Ibagué')
        cls.o102 = Obra.objects.create(constructora=cls.asf, nombre='iba10-2', codigo_obra='IBA10-611')
        Obra.objects.create(constructora=cls.asf, nombre='IBA 10-1', codigo_obra='IBA10-612')

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        super().tearDownClass()

    def test_por_nombre_de_obra_con_codigo(self):
        obra, motivo, _ = resolver_obra('Ibagué', '8-1', self.lectura)
        self.assertEqual(obra, self.o81)
        self.assertIn('nombre de obra', motivo)
        obra, motivo, _ = resolver_obra('Ibagué', '10-2', self.lectura)
        self.assertEqual(obra, self.o102)

    def test_unica_obra_solo_si_el_excel_tiene_una_sola_columna(self):
        Obra.objects.filter(pk=self.o81.pk).update(nombre='Bodega')
        obra, motivo, _ = resolver_obra('Ibagué', '8-1', self.lectura, columnas_de_la_empresa=1)
        self.assertEqual(obra, self.o81)
        self.assertIn('supuesto', motivo)
        obra, motivo, _ = resolver_obra('Ibagué', '8-1', self.lectura, columnas_de_la_empresa=9)
        self.assertIsNone(obra)
        self.assertIn('--crear-obras', motivo)

    def test_dos_columnas_no_caen_en_la_misma_obra(self):
        # empresa 10 tiene 2 columnas en un libro con '10-2' y '10-3'; '10-3' no existe -> sin resolver, no a IBA 10-2
        obra, motivo, _ = resolver_obra('Ibagué', '10-3', self.lectura, columnas_de_la_empresa=2)
        self.assertIsNone(obra)

    def test_crear_obras(self):
        out = StringIO()
        call_command('importar_catalogo_excel', self.ruta, '--ciudad', 'Ibagué', '--crear-obras', stdout=out)
        # 8-1 y 10-2 ya existían: no crea nada
        self.assertNotIn('obra(s) creadas', out.getvalue())
        Obra.objects.filter(pk=self.o81.pk).delete()
        out = StringIO()
        call_command('importar_catalogo_excel', self.ruta, '--ciudad', 'Ibagué', '--crear-obras', stdout=out)
        self.assertIn('1 obra(s) creadas', out.getvalue())
        nueva = Obra.objects.get(constructora=self.sika, nombre='AMBALA (IBA 8-1)')
        self.assertEqual(nueva.codigo_obra, 'IBA8-1')
        self.assertTrue(PrecioServicio.objects.filter(obra=nueva).exists())
        # sin --crear-obras y sin obra: queda sin resolver, no se inventa nada
        out = StringIO()
        Obra.objects.filter(pk=nueva.pk).delete()
        call_command('importar_catalogo_excel', self.ruta, '--ciudad', 'Ibagué', '--dry-run', stdout=out)
        self.assertIn('no tiene obras en el sistema', out.getvalue())


class ResolverObraCodigoMandaTests(TestCase):
    """En WordPress hay razones sociales mal capturadas: el código de empresa manda sobre la razón social."""

    @classmethod
    def setUpTestData(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = str(Path(cls.tmp.name) / 'base.xlsx')
        _libro_de_prueba(cls.ruta)          # empresa 8 = 'DYCO SAS', empresa 10 = 'ASFALTEMOS'
        cls.lectura = leer_lista_precios(cls.ruta)
        # IBA10 existe pero con OTRA razón social; una constructora ajena se llama 'ASFALTEMOS'
        cls.iba10 = Constructora.objects.create(codigo='IBA10', nombre='AP COELLO', ciudad='Ibagué')
        cls.o = Obra.objects.create(constructora=cls.iba10, nombre='IBA 10-2', codigo_obra='IBA10-9')
        ajena = Constructora.objects.create(codigo='IBA14', nombre='ASFALTEMOS', ciudad='Ibagué')
        Obra.objects.create(constructora=ajena, nombre='IBA 14-1', codigo_obra='IBA14-1')

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        super().tearDownClass()

    def test_codigo_gana_a_razon_social(self):
        obra, motivo, _ = resolver_obra('Ibagué', '10-2', self.lectura)
        self.assertEqual(obra, self.o)

    def test_razon_social_solo_si_no_hay_codigo(self):
        Constructora.objects.filter(codigo='IBA10').update(codigo='IBA77')
        # sin código IBA10, 'ASFALTEMOS' lleva a IBA14; con 2 columnas en el Excel no aplica "única obra"
        obra, motivo, candidatas = resolver_obra('Ibagué', '10-2', self.lectura, columnas_de_la_empresa=2)
        self.assertIsNone(obra)
        self.assertEqual([c.codigo_obra for c in candidatas], ['IBA14-1'])

    def test_sufijo_de_codigo_reutilizado_cuenta(self):
        Constructora.objects.filter(codigo='IBA10').update(codigo='IBA10B')
        obra, motivo, _ = resolver_obra('Ibagué', '10-2', self.lectura)
        self.assertEqual(obra, self.o)


class IntegridadFacturacionTests(TestCase):
    """Doble clic, conexiones cortadas y cambios cruzados: nada se duplica ni se descuadra."""

    @classmethod
    def setUpTestData(cls):
        cls.staff = UsuarioBase.objects.create_user('admin_int', password='clave-de-prueba', es_geolab=True)
        FuncionarioGeolab.objects.create(user=cls.staff, area='admin')
        cls.c = Constructora.objects.create(nombre='Buc SAS', codigo='B1', ciudad='Bucaramanga')
        cls.obra = Obra.objects.create(nombre='Obra Buc', codigo_obra='B1-1', constructora=cls.c)
        cat = CategoriaServicio.objects.create(ciudad='Bucaramanga', codigo='1', nombre='CONCRETOS')
        cls.srv = TipoServicio.objects.create(categoria=cat, codigo='1-1', nombre='Compresión')
        cls.srv2 = TipoServicio.objects.create(categoria=cat, codigo='1-2', nombre='Flexión')
        PrecioServicio.objects.create(obra=cls.obra, tipo_servicio=cls.srv, precio=1000)
        PrecioServicio.objects.create(obra=cls.obra, tipo_servicio=cls.srv2, precio=None)
        Impuesto.objects.create(nombre='IVA', porcentaje='19.00')

    def setUp(self):
        self.client.login(username='admin_int', password='clave-de-prueba')
        session = self.client.session
        session[SESSION_KEY] = 'Bucaramanga'
        session.save()
        self.pdf = mock.patch('facturacion.views._generar_pdf_factura')
        self.pdf.start()
        self.addCleanup(self.pdf.stop)

    def _registro(self, **extra):
        datos = dict(obra=self.obra, tipo_servicio=self.srv, fecha_realizacion='2026-09-10',
                     cantidad=2, numero_informe='INF-1', precio_unitario_congelado=1000, creado_por=self.staff)
        datos.update(extra)
        return RegistroServicio.objects.create(**datos)

    def _finalizar(self, ids):
        return self.client.post(reverse('finalizar_factura'), {
            'obra_id': self.obra.pk, 'fecha_inicio': '2026-09-01', 'fecha_fin': '2026-09-30',
            'registro_ids': [str(i) for i in ids],
        })

    def test_doble_envio_no_crea_dos_facturas(self):
        r = self._registro()
        r1 = self._finalizar([r.pk])
        self.assertEqual(Factura.objects.count(), 1)
        self.assertRedirects(r1, reverse('repositorio_facturas_obra', args=[self.obra.pk]), fetch_redirect_response=False)
        r2 = self._finalizar([r.pk])    # el segundo clic llega con los mismos ids
        self.assertEqual(Factura.objects.count(), 1)
        self.assertRedirects(r2, reverse('repositorio_facturas_obra', args=[self.obra.pk]), fetch_redirect_response=False)
        f = Factura.objects.get()
        self.assertEqual((f.registros.count(), f.subtotal, f.monto_iva, f.total), (1, 2000, 380, 2380))

    def test_vista_previa_desactualizada_no_factura(self):
        a = self._registro()
        b = self._registro(numero_informe='INF-2')    # registrado después de la vista previa
        r = self._finalizar([a.pk])
        self.assertEqual(Factura.objects.count(), 0)
        self.assertIn('fecha_inicio=2026-09-01', r.url)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertIsNone(a.factura_id)
        self.assertIsNone(b.factura_id)

    def test_anular_pagada_requiere_volver_a_pendiente(self):
        r = self._registro()
        self._finalizar([r.pk])
        f = Factura.objects.get()
        self.client.post(reverse('marcar_factura_pagada', args=[f.pk]))
        self.client.post(reverse('anular_factura', args=[f.pk]))
        f.refresh_from_db()
        self.assertEqual(f.estado, 'PAGADA')
        self.assertEqual(f.registros.count(), 1)
        self.client.post(reverse('marcar_factura_pendiente', args=[f.pk]))
        self.client.post(reverse('anular_factura', args=[f.pk]))
        f.refresh_from_db(); r.refresh_from_db()
        self.assertEqual(f.estado, 'ANULADA')
        self.assertIsNone(r.factura_id)

    def test_editar_conserva_precio_congelado_y_no_pone_cero(self):
        r = self._registro()
        PrecioServicio.objects.filter(obra=self.obra, tipo_servicio=self.srv).update(precio=5000)
        datos = {'obra': self.obra.pk, 'tipo_servicio': self.srv.pk, 'fecha_realizacion': '2026-09-10',
                 'numero_informe': 'INF-1 corregido', 'cantidad': 3}
        self.client.post(reverse('editar_registro_facturacion', args=[r.pk]), datos)
        r.refresh_from_db()
        self.assertEqual((r.precio_unitario_congelado, r.cantidad, r.subtotal), (1000, 3, 3000))
        # cambiar a un servicio sin precio en la obra: no se guarda
        datos['tipo_servicio'] = self.srv2.pk
        resp = self.client.post(reverse('editar_registro_facturacion', args=[r.pk]), datos)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'no tiene precio')
        r.refresh_from_db()
        self.assertEqual((r.tipo_servicio_id, r.precio_unitario_congelado), (self.srv.pk, 1000))

    def _crear(self, servicios, envio_id='11111111-2222-4333-8444-555555555555', fecha='2026-09-10'):
        return self.client.post(reverse('crear_registro_facturacion'), data=json.dumps({
            'obra_id': self.obra.pk, 'fecha_realizacion': fecha, 'servicios': servicios, 'envio_id': envio_id,
        }), content_type='application/json')

    def test_registro_reintentado_no_se_duplica_y_es_todo_o_nada(self):
        linea = {'tipo_servicio_id': self.srv.pk, 'numero_informe': 'INF-9', 'cantidad': 2}
        r = self._crear([linea])
        self.assertEqual(r.json()['status'], 'success')
        self.assertEqual(RegistroServicio.objects.count(), 1)
        r = self._crear([linea])    # se cortó la conexión y el usuario volvió a dar Guardar
        self.assertTrue(r.json()['repetido'])
        self.assertEqual(RegistroServicio.objects.count(), 1)
        # una línea mala (cantidad 0) tumba todo el envío
        r = self._crear([dict(linea, numero_informe='INF-10'), dict(linea, cantidad=0)], envio_id=None)
        self.assertEqual(r.status_code, 400)
        self.assertIn('cantidad debe ser 1', r.json()['message'])
        self.assertEqual(RegistroServicio.objects.count(), 1)
        r = self._crear([linea], envio_id=None, fecha='2099-01-01')
        self.assertEqual(r.status_code, 400)
        # sesión vencida: JSON, no la página de login
        self.client.logout()
        r = self._crear([linea])
        self.assertEqual(r.status_code, 401)
        self.assertTrue(r.json()['sesion'])

    def test_seed_no_corre_en_produccion(self):
        from django.core.management import call_command
        from django.core.management.base import CommandError
        from django.test import override_settings
        with override_settings(DEBUG=False), self.assertRaises(CommandError):
            call_command('seed_facturacion')

    def test_iva_duplicado_o_ausente_no_rompe_la_vista_previa(self):
        self._registro()
        Impuesto.objects.create(nombre='iva', porcentaje='5.00')
        r = self.client.get(reverse('generar_factura'), {'obra': self.obra.pk, 'fecha_inicio': '2026-09-01', 'fecha_fin': '2026-09-30'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['iva_porcentaje'], Decimal('19.00'))
        Impuesto.objects.update(activo=False)
        r = self.client.get(reverse('generar_factura'), {'obra': self.obra.pk, 'fecha_inicio': '2026-09-01', 'fecha_fin': '2026-09-30'})
        self.assertFalse(r.context['hay_iva'])
        self.assertContains(r, 'sin IVA')

    def test_api_precio_cero_y_otra_sede(self):
        PrecioServicio.objects.filter(obra=self.obra, tipo_servicio=self.srv).update(precio=0)
        r = self.client.get(reverse('api_get_precio'), {'obra_id': self.obra.pk, 'tiposervicio_id': self.srv.pk})
        self.assertEqual(r.json()['precio'], '0.00')
        session = self.client.session
        session[SESSION_KEY] = 'Bogotá'
        session.save()
        r = self.client.get(reverse('api_get_precio'), {'obra_id': self.obra.pk, 'tiposervicio_id': self.srv.pk})
        self.assertIsNone(r.json()['precio'])

    def test_no_cambiar_ciudad_de_empresa_con_facturas_ni_categoria_7_con_registros(self):
        from core.forms import ConstructoraForm
        from .forms import CategoriaServicioForm
        self._registro()
        form = ConstructoraForm({'nombre': 'Buc SAS', 'ciudad': 'Bogotá', 'numero_codigo': '1', 'nit': ''}, instance=self.c)
        self.assertFalse(form.is_valid())
        self.assertIn('ciudad', form.errors)
        cat = self.srv.categoria
        form = CategoriaServicioForm({'codigo': '7', 'nombre': 'CONCRETOS'}, instance=cat)
        self.assertFalse(form.is_valid())
        self.assertIn('transporte', form.errors['codigo'][0])

